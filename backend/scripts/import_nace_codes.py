"""'Risk kart.xlsx' dosyasındaki 'nace kodları' sayfasını data/nace_kodlari.json'a aktarır.

Kullanım: python scripts/import_nace_codes.py /path/to/Risk\\ kart.xlsx
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

GROUP_RE = re.compile(r'^\d+\.\s*GRUP AD[Iİ]\s*:\s*(.+)$', re.IGNORECASE)
CODE_RE = re.compile(r'^(\d{2}(?:\.\d{2}){1,2})\s+(.+)$')

DEST = Path(__file__).resolve().parent.parent / "data" / "nace_kodlari.json"


def main(src_path: str):
    wb = openpyxl.load_workbook(src_path, data_only=True)
    ws = wb["nace kodları"]

    records = []
    current_group = None
    for r in range(2, ws.max_row + 1):
        val = ws.cell(row=r, column=1).value
        if not val:
            continue
        val = val.strip()

        group_match = GROUP_RE.match(val)
        if group_match:
            current_group = group_match.group(1).strip().rstrip("\xa0").strip()
            continue

        code_match = CODE_RE.match(val)
        if code_match:
            records.append({
                "kod": code_match.group(1),
                "aciklama": code_match.group(2).strip(),
                "grup": current_group,
            })

    DEST.write_text(json.dumps(records, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"{len(records)} NACE kodu yazıldı: {DEST}")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("Kullanım: python scripts/import_nace_codes.py /path/to/Risk_kart.xlsx")
        sys.exit(1)
    main(sys.argv[1])
